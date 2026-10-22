import docx
import re

d = docx.Document(r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0\.user_uploaded\media_1792596558882_0f1ebc8d.docx")

blockchain_terms = [
    'ChainGuard Sandbox', 'ChainGuard', 'smart contract', 'Solidity', 'Hardhat',
    'solc-js', 'solc', 'reentrancy', 'Reentrancy', 'gas profiling', 'gas usage',
    'gas consumption', 'Ethers.js', 'vulnerability', 'attack simulation', 'overflow'
]

count = 0
for i, p in enumerate(d.paragraphs):
    for term in blockchain_terms:
        if term.lower() in p.text.lower():
            count += 1
            break

print(f"Total paragraphs with blockchain terms: {count} out of {len(d.paragraphs)}")

tbl_count = 0
for tidx, tbl in enumerate(d.tables):
    has_term = False
    for row in tbl.rows:
        for c in row.cells:
            for term in blockchain_terms:
                if term.lower() in c.text.lower():
                    has_term = True
                    break
            if has_term: break
        if has_term: break
    if has_term:
        tbl_count += 1

print(f"Total tables with blockchain terms: {tbl_count} out of {len(d.tables)}")
