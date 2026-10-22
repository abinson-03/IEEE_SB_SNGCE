import os
import re
import subprocess
import markdown
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

MD_PATH = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.md"
DOCX_OUT = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.docx"
PDF_OUT = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.pdf"
HTML_TEMP = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\report_styled.html"

DOWNLOADS_DIR = r"C:\Users\Abinson Babu\Downloads\ieee-sb-portal\docs"
ARTIFACT_DIR = r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0"

CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
if not os.path.exists(CHROME_PATH):
    CHROME_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

def set_cell_border(cell, **kwargs):
    """
    Set cell borders
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='CBD5E1', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def build_docx(md_text, docx_path):
    print(f"Building DOCX: {docx_path}...")
    doc = docx.Document()

    # Set page margins: KTU standard (Left 1.25", Right 1.0", Top 1.0", Bottom 1.0")
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27) # A4 width
        section.page_height = Inches(11.69) # A4 height

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    normal_style.paragraph_format.line_spacing = 1.18
    normal_style.paragraph_format.space_after = Pt(6)

    lines = md_text.splitlines()
    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []

    page_break_keywords = [
        "DECLARATION", "CERTIFICATE", "ACKNOWLEDGEMENT", "ABSTRACT",
        "CONTENTS", "LIST OF TABLES", "LIST OF FIGURES", "ABBREVIATIONS",
        "NOTATIONS", "SCRUM PROJECT FRAMEWORK", "CHAPTER 1.", "CHAPTER 2.",
        "CHAPTER 3.", "CHAPTER 4.", "CHAPTER 5.", "CHAPTER 6.", "CHAPTER 7.",
        "CHAPTER 8.", "CHAPTER 9.", "CHAPTER 10.", "REFERENCES"
    ]

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # Handle Code / Mermaid blocks
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Courier New'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(lines[i])
            i += 1
            continue

        # Handle Markdown Tables
        if line.startswith("|") and line.endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            # Process collected table
            process_table(doc, table_lines)
            in_table = False
            table_lines = []

        if not line:
            i += 1
            continue

        # Check for section breaks
        upper_line = line.upper()
        if any(kw in upper_line for kw in page_break_keywords) and (line.startswith("#") or len(line) < 50):
            # Add page break before major section
            if len(doc.paragraphs) > 5:
                doc.add_page_break()

        # Process Headings
        if line.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_md_inline(line[2:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x49, 0x76) # IEEE Deep Blue
        elif line.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_md_inline(line[3:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x49, 0x76)
        elif line.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_md_inline(line[4:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif line.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_md_inline(line[5:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            num_match = re.match(r"^(\d+\.)\s*(.*)", line)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, num_match.group(2))
        elif line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            clean_q = re.sub(r"^>\s*", "", line)
            add_formatted_runs(p, clean_q, italic=True)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            # Center title page lines
            if "DEPARTMENT OF COMPUTER APPLICATIONS" in upper_line or "PROJECT REPORT" in upper_line or "SREE NARAYANA GURUKULAM" in upper_line or "NOVEMBER 2026" in upper_line:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_runs(p, line)

        i += 1

    # In case document ended with a table
    if in_table and table_lines:
        process_table(doc, table_lines)

    doc.save(docx_path)
    print(f"DOCX saved successfully to {docx_path}")

def clean_md_inline(text):
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()

def add_formatted_runs(paragraph, text, italic=False):
    # Regex for bold, code, regular segments
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
            run.font.name = 'Times New Roman'
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x00, 0x49, 0x76)
        elif tok.startswith("*") and tok.endswith("*"):
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
            run.font.name = 'Times New Roman'
        else:
            run = paragraph.add_run(tok)
            run.font.name = 'Times New Roman'
        if italic:
            run.italic = True

def process_table(doc, table_lines):
    rows = []
    for line in table_lines:
        parts = [c.strip() for c in line.split("|")[1:-1]]
        # Ignore separator lines like |:---|:---|
        if parts and all(re.match(r"^:?-+:?$", p) for p in parts if p):
            continue
        if parts:
            rows.append(parts)

    if not rows:
        return

    num_cols = max(len(r) for r in rows)
    num_rows = len(rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    border_spec = dict(sz=4, val='single', color='CBD5E1', space='0')

    for r_idx, row_data in enumerate(rows):
        is_header = (r_idx == 0)
        row = table.rows[r_idx]
        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05

            add_formatted_runs(p, val)
            if is_header:
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
                set_cell_background(cell, "004976") # IEEE Deep Navy header
            else:
                if r_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC") # alternating row shading
                for r in p.runs:
                    r.font.size = Pt(10)

            set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)

    # Space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(8)

def build_pdf_from_html(md_text, html_temp, pdf_out):
    print(f"Converting markdown to styled HTML: {html_temp}...")
    html_content = markdown.markdown(md_text, extensions=['tables', 'fenced_code', 'nl2br'])

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>IEEE SB SNGCE Portal - Project Report</title>
<style>
  @page {{
    size: A4;
    margin: 25mm 20mm 20mm 25mm;
    @bottom-center {{
      content: counter(page);
      font-family: "Times New Roman", serif;
      font-size: 10pt;
      color: #64748b;
    }}
  }}
  *, *:before, *:after {{
    box-sizing: border-box;
  }}
  body {{
    font-family: "Times New Roman", Times, Georgia, serif;
    font-size: 12pt;
    line-height: 1.5;
    color: #111827;
    background: #ffffff;
    margin: 0;
    padding: 0;
  }}
  h1 {{
    font-size: 18pt;
    color: #004976;
    text-transform: uppercase;
    border-bottom: 2px solid #004976;
    padding-bottom: 4px;
    margin-top: 24pt;
    margin-bottom: 12pt;
    page-break-before: always;
  }}
  h1:first-of-type {{
    page-break-before: avoid;
    text-align: center;
    border-bottom: none;
  }}
  h2 {{
    font-size: 14pt;
    color: #004976;
    margin-top: 18pt;
    margin-bottom: 8pt;
    page-break-after: avoid;
  }}
  h3 {{
    font-size: 12.5pt;
    color: #0f172a;
    margin-top: 14pt;
    margin-bottom: 6pt;
    page-break-after: avoid;
  }}
  h4 {{
    font-size: 12pt;
    color: #334155;
    margin-top: 10pt;
    margin-bottom: 4pt;
    page-break-after: avoid;
  }}
  p {{
    margin: 0 0 8pt 0;
    text-align: justify;
  }}
  ul, ol {{
    margin: 4pt 0 10pt 20pt;
    padding: 0;
  }}
  li {{
    margin-bottom: 4pt;
    text-align: justify;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 12pt 0;
    font-size: 10pt;
    page-break-inside: avoid;
  }}
  th {{
    background-color: #004976;
    color: #ffffff;
    font-weight: bold;
    border: 1px solid #94a3b8;
    padding: 7px 10px;
    text-align: left;
  }}
  td {{
    border: 1px solid #cbd5e1;
    padding: 6px 10px;
    vertical-align: top;
  }}
  tr:nth-child(even) td {{
    background-color: #f8fafc;
  }}
  pre, code {{
    font-family: "Courier New", Courier, monospace;
  }}
  code {{
    background: #f1f5f9;
    color: #004976;
    padding: 2px 5px;
    border-radius: 3px;
    font-size: 10pt;
  }}
  pre {{
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-left: 4px solid #004976;
    padding: 10px 14px;
    font-size: 9.5pt;
    overflow-x: auto;
    page-break-inside: avoid;
  }}
  blockquote {{
    border-left: 4px solid #004976;
    background: #f0f7ff;
    margin: 10pt 0;
    padding: 8pt 14pt;
    font-style: italic;
  }}
  hr {{
    border: none;
    border-top: 1px solid #cbd5e1;
    margin: 18pt 0;
  }}
  .title-center {{
    text-align: center;
  }}
</style>
</head>
<body>
{html_content}
</body>
</html>
"""
    with open(html_temp, "w", encoding="utf-8") as f:
        f.write(full_html)

    print(f"Generating PDF using headless browser: {CHROME_PATH}...")
    cmd = [
        CHROME_PATH,
        "--headless=new",
        "--disable-gpu",
        f"--print-to-pdf={pdf_out}",
        "--no-pdf-header-footer",
        html_temp
    ]
    subprocess.run(cmd, check=True)
    print(f"PDF successfully generated at {pdf_out}!")

def main():
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    # 1. Build DOCX
    build_docx(md_text, DOCX_OUT)

    # 2. Build PDF
    build_pdf_from_html(md_text, HTML_TEMP, PDF_OUT)

    # 3. Synchronize to Downloads and Artifact directory
    for src, fname in [(DOCX_OUT, "PROJECT_REPORT.docx"), (PDF_OUT, "PROJECT_REPORT.pdf")]:
        dst_dl = os.path.join(DOWNLOADS_DIR, fname)
        dst_art = os.path.join(ARTIFACT_DIR, fname)
        with open(src, "rb") as sf:
            b = sf.read()
            with open(dst_dl, "wb") as df:
                df.write(b)
            with open(dst_art, "wb") as af:
                af.write(b)
        print(f"Synced {fname} to Downloads and Artifact Directory.")

    print("\nALL DOCUMENTS GENERATED AND SYNCED SUCCESSFULLY! ✓")

if __name__ == "__main__":
    main()
