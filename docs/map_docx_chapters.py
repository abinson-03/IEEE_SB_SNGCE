import docx

doc = docx.Document(r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0\.user_uploaded\media_1792596558882_0f1ebc8d.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

chapters = []
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.upper().startswith("CHAPTER ") or t.upper().startswith("1.") or t.upper().startswith("2.") or t.upper().startswith("3.") or t.upper().startswith("4.") or t.upper().startswith("5.") or t.upper().startswith("6.") or t.upper().startswith("7.") or t.upper().startswith("8.") or t.upper().startswith("9.") or t.upper().startswith("10."):
        if len(t) < 60:
            chapters.append((i, t))

for c in chapters[:30]:
    print(f"[{c[0]}] {c[1]}")
