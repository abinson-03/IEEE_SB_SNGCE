import os
import re
import shutil
import subprocess
import markdown
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

MD_PATH = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.md"
DOCX_OUT = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.docx"
PDF_OUT = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.pdf"
HTML_TEMP = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\report_clean_styled.html"

DOWNLOADS_DIR = r"C:\Users\Abinson Babu\Downloads\ieee-sb-portal\docs"
ARTIFACT_DIR = r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0"
FIGURES_DIR = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\figures"
LOGO_PATH = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\build_abinson\unpacked\word\media\image1.jpg"

CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
if not os.path.exists(CHROME_PATH):
    CHROME_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = f'w:{edge}'
        elm = tcBorders.find(qn(tag))
        if elm is None:
            elm = OxmlElement(tag)
            tcBorders.append(elm)
        elm.set(qn('w:val'), 'single')
        elm.set(qn('w:sz'), '4')
        elm.set(qn('w:space'), '0')
        elm.set(qn('w:color'), 'CBD5E1')

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def clean_md_inline(text):
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("&nbsp;", " ")
    return text.strip()

def add_formatted_runs(paragraph, text, italic=False):
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
            run.font.name = 'Times New Roman'
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x00, 0x49, 0x76)
        elif tok.startswith("*") and tok.endswith("*"):
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
            run.font.name = 'Times New Roman'
        else:
            run = paragraph.add_run(tok.replace("&nbsp;", " "))
            run.font.name = 'Times New Roman'
        if italic:
            run.italic = True

def process_table(doc, table_lines):
    rows_data = []
    for line in table_lines:
        parts = [c.strip() for c in line.split("|")[1:-1]]
        if parts and all(re.match(r"^:?-+:?$", p) for p in parts if p):
            continue
        if parts:
            rows_data.append(parts)

    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for r_idx, r_data in enumerate(rows_data):
        is_header = (r_idx == 0)
        row = table.rows[r_idx]
        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell_text = r_data[c_idx] if c_idx < len(r_data) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05

            add_formatted_runs(p, cell_text)
            set_cell_borders(cell)

            if is_header:
                set_cell_background(cell, "004976")
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.5)
            else:
                if r_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.name = 'Times New Roman'

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_docx_report():
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    print(f"Building complete authentic DOCX: {DOCX_OUT}...")
    doc = docx.Document()

    # KTU standard margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.25)
        sec.right_margin = Inches(1.0)
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    lines = md_text.splitlines()
    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []

    page_break_keywords = [
        "DECLARATION", "CERTIFICATE", "ACKNOWLEDGEMENT", "ABSTRACT",
        "CONTENTS", "LIST OF TABLES", "LIST OF FIGURES", "ABBREVIATIONS",
        "NOTATIONS", "SCRUM PROJECT FRAMEWORK", "CHAPTER 1.", "CHAPTER 2.",
        "CHAPTER 3.", "CHAPTER 4.", "CHAPTER 5.", "CHAPTER 6.", "CHAPTER 7.",
        "CHAPTER 8.", "CHAPTER 9.", "CHAPTER 10.", "REFERENCES"
    ]

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # College Logo Placeholder
        if line == "[COLLEGE_LOGO_HERE]":
            if os.path.exists(LOGO_PATH):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run()
                run.add_picture(LOGO_PATH, width=Inches(1.8))
            i += 1
            continue

        # Dynamic Figure Insertion Tag
        if line.startswith("[INSERT_FIGURE:"):
            m = re.match(r"\[INSERT_FIGURE:\s*([^|]+)\|\s*([^\]]+)\]", line)
            if m:
                fig_filename = m.group(1).strip()
                fig_caption = m.group(2).strip()
                fig_path = os.path.join(FIGURES_DIR, fig_filename)
                if os.path.exists(fig_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(12)
                    p_img.paragraph_format.space_after = Pt(4)
                    run_img = p_img.add_run()
                    # Add image with exact 5.8 inches width - python-docx auto-calculates height proportionally!
                    run_img.add_picture(fig_path, width=Inches(5.8))

                    # Caption
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(12)
                    run_cap = p_cap.add_run(fig_caption)
                    run_cap.font.name = 'Times New Roman'
                    run_cap.font.size = Pt(10)
                    run_cap.font.bold = True
                    run_cap.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                    print(f"  Inserted figure: {fig_filename} with caption: {fig_caption[:50]}...")
            i += 1
            continue

        # Code blocks
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.right_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Courier New'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(lines[i])
            i += 1
            continue

        # Tables
        if line.startswith("|") and line.endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            process_table(doc, table_lines)
            in_table = False
            table_lines = []

        if not line:
            i += 1
            continue

        # Section breaks
        upper_line = line.upper()
        if any(kw in upper_line for kw in page_break_keywords) and (line.startswith("#") or len(line) < 45):
            if len(doc.paragraphs) > 5:
                doc.add_page_break()

        # Headings
        if line.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = h.add_run(clean_md_inline(line[2:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x49, 0x76)
        elif line.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_md_inline(line[3:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x49, 0x76)
        elif line.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_md_inline(line[4:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif line.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_md_inline(line[5:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            num_match = re.match(r"^(\d+\.)\s*(.*)", line)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, num_match.group(2))
        elif line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            clean_q = re.sub(r"^>\s*", "", line)
            add_formatted_runs(p, clean_q, italic=True)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            if any(k in upper_line for k in ["DEPARTMENT OF COMPUTER APPLICATIONS", "PROJECT REPORT", "SREE NARAYANA GURUKULAM", "NOVEMBER 2026", "SUBMITTED BY", "UNDER THE GUIDANCE OF"]):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_runs(p, line)

        i += 1

    if in_table and table_lines:
        process_table(doc, table_lines)

    doc.save(DOCX_OUT)
    print(f"DOCX saved successfully to {DOCX_OUT}!")

    # Deliver to Downloads & Artifacts
    for dest_dir in [DOWNLOADS_DIR, ARTIFACT_DIR]:
        target = os.path.join(dest_dir, "PROJECT_REPORT.docx")
        try:
            shutil.copy2(DOCX_OUT, target)
            print(f"Copied to {target}")
        except Exception as e:
            alt = os.path.join(dest_dir, "PROJECT_REPORT_FINAL_SUBMISSION.docx")
            shutil.copy2(DOCX_OUT, alt)
            print(f"Copied to alternate: {alt}")

build_docx_report()
