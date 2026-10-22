import docx

doc = docx.Document(r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\build_abinson\stage2.docx")

print("=== Paragraphs 230 to 260 ===")
for i in range(230, 260):
    t = doc.paragraphs[i].text.strip()
    if t:
        print(f"[{i}] {t}")

print("\n=== Paragraphs 1180 to 1205 ===")
for i in range(1180, 1205):
    t = doc.paragraphs[i].text.strip()
    if t:
        print(f"[{i}] {t}")
