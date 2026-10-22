import docx
import re

d = docx.Document(r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0\.user_uploaded\media_1792596558882_0f1ebc8d.docx")

rels = d.part.rels

for i, p in enumerate(d.paragraphs):
    xml = p._p.xml
    m = re.findall(r'r:embed="([^"]+)"', xml)
    if m:
        caption = ''
        for offset in range(-2, 3):
            if 0 <= i + offset < len(d.paragraphs):
                t = d.paragraphs[i + offset].text.strip()
                if any(k in t for k in ['Figure', 'Fig', 'Wireframe', 'Diagram', 'Chart', 'Architecture', 'Structure']):
                    caption = t
        print(f"P[{i}]: {m} -> {[rels[r].target_ref for r in m if r in rels]} | {caption[:80]}")
