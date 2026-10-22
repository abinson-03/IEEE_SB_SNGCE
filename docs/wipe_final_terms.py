import docx

doc = docx.Document(r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.docx")

for i, p in enumerate(doc.paragraphs):
    if "blockchain" in p.text.lower():
        p.text = p.text.replace("blockchain technology", "full-stack web technology")
        p.text = p.text.replace("Blockchain technology", "Full-stack web technology")
        p.text = p.text.replace("blockchain network", "web application platform")
        p.text = p.text.replace("blockchain environment", "cloud-ready web environment")
        p.text = p.text.replace("blockchain", "web portal")
        p.text = p.text.replace("Blockchain", "Web portal")

for i, p in enumerate(doc.paragraphs):
    if "solidity" in p.text.lower():
        p.text = p.text.replace("Solidity", "MERN Stack").replace("solidity", "MERN stack")

for ti, t in enumerate(doc.tables):
    for ri, r in enumerate(t.rows):
        for ci, c in enumerate(r.cells):
            if "blockchain" in c.text.lower() or "solidity" in c.text.lower():
                c.text = c.text.replace("blockchain technology", "full-stack web technology")
                c.text = c.text.replace("Blockchain technology", "Full-stack web technology")
                c.text = c.text.replace("blockchain network", "web application platform")
                c.text = c.text.replace("blockchain environment", "cloud-ready web environment")
                c.text = c.text.replace("blockchain", "portal")
                c.text = c.text.replace("Blockchain", "Portal")
                c.text = c.text.replace("Solidity", "MERN Stack").replace("solidity", "MERN stack")

doc.save(r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.docx")
doc.save(r"C:\Users\Abinson Babu\Downloads\ieee-sb-portal\docs\PROJECT_REPORT_IEEE_PORTAL.docx")
doc.save(r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0\PROJECT_REPORT.docx")

print("--- AUDIT RESULTS ---")
for t in ['blockchain', 'solidity']:
    p_cnt = sum(1 for p in doc.paragraphs if t in p.text.lower())
    t_cnt = sum(1 for tbl in doc.tables for r in tbl.rows for c in r.cells if t in c.text.lower())
    print(f"{t}: {p_cnt} in paragraphs, {t_cnt} in tables")
