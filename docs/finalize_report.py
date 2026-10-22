import os
import shutil
import docx

DOCX_IN = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.docx"
DOCX_OUT = r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\PROJECT_REPORT.docx"
DOWNLOADS_ALT = r"C:\Users\Abinson Babu\Downloads\ieee-sb-portal\docs\PROJECT_REPORT_PERFECT.docx"
ARTIFACT_OUT = r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0\PROJECT_REPORT.docx"

doc = docx.Document(DOCX_IN)

# 1. Fix Paragraph 2010 and 2024 (CHAPTER 7 -> CHAPTER 8)
p2010 = doc.paragraphs[2010]
p2010.text = p2010.text.replace("CHAPTER 7", "CHAPTER 8")
print("Fixed [2010]:", p2010.text.strip())

p2024 = doc.paragraphs[2024]
p2024.text = p2024.text.replace("CHAPTER 7", "CHAPTER 8")
print("Fixed [2024]:", p2024.text.strip())

# 2. Fix the 1 remaining sandbox paragraph [1298]
for p in doc.paragraphs:
    if "sandbox" in p.text.lower():
        p.text = p.text.replace("Sandbox", "Portal").replace("sandbox", "portal")
        print("Fixed sandbox in paragraph:", p.text[:80])

# 3. Fix Table 25 (Planning Poker)
tbl25 = doc.tables[25]
new_table_data = [
    ["User Story ID", "Abinson (Dev)", "Complexity", "Effort (Hrs)", "Final Estimate", "Feature Scope & Acceptance Criteria", "User Story ID"],
    ["US02", "2", "Low", "3", "2", "Event and chapter membership management", "US02"],
    ["US03", "3", "Medium", "5", "3", "Interactive chapter portfolio and event schedule", "US03"],
    ["US05", "5", "High", "6", "5", "Dynamic QR attendance check-in workflow", "US05"],
    ["US06", "5", "High", "7", "5", "Secure event registration with UPI payment validation", "US06"],
    ["US07", "8", "Very High", "10", "8", "Automated PDFKit certificate generation and email dispatch", "US07"],
]

for row_idx, row_data in enumerate(new_table_data):
    row = tbl25.rows[row_idx]
    for col_idx, text_val in enumerate(row_data):
        row.cells[col_idx].text = text_val

print("Table 25 updated successfully.")

# 4. Check Table 29 (Story Point Summary for Abinson Babu)
tbl29 = doc.tables[29]
tbl29_data = [
    ["Functional Role / Domain", "Assigned Stories", "Story Points", "Estimated Hours"],
    ["Abinson Babu (Backend & API Engine)", "6", "21", "28"],
    ["Abinson Babu (Frontend & UI Engineering)", "6", "20", "26"],
    ["Abinson Babu (Security, Testing & DevOps)", "4", "15", "20"],
    ["Total (Individual Project)", "16", "56", "74"],
]
for row_idx, row_data in enumerate(tbl29_data):
    if row_idx < len(tbl29.rows):
        row = tbl29.rows[row_idx]
        for col_idx, text_val in enumerate(row_data):
            if col_idx < len(row.cells):
                row.cells[col_idx].text = text_val

print("Table 29 updated successfully.")

# Save updated docx
doc.save(DOCX_OUT)
print("Saved final DOCX to:", DOCX_OUT)

# Deliver to destinations
shutil.copy2(DOCX_OUT, DOWNLOADS_ALT)
print("Delivered to Downloads:", DOWNLOADS_ALT)

shutil.copy2(DOCX_OUT, ARTIFACT_OUT)
print("Delivered to Artifact:", ARTIFACT_OUT)
