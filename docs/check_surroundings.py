import docx

d = docx.Document(r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0\.user_uploaded\media_1792596558882_0f1ebc8d.docx")

indices = [367, 417, 434, 564, 570, 1160, 1303, 1324, 1358, 1865]
for idx in indices:
    print(f"=== Paragraph {idx} ===")
    start = max(0, idx - 3)
    end = min(len(d.paragraphs), idx + 4)
    for j in range(start, end):
        txt = d.paragraphs[j].text.strip().encode('ascii', errors='replace').decode()
        if txt:
            print(f"  [{j}] {txt[:90]}")
