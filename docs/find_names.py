import docx

doc = docx.Document(r"C:\Users\Abinson Babu\Desktop\ieee-sb-portal\docs\build_abinson\stage2.docx")

print("=== PARAGRAPHS WITH NAMES ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if any(n in txt.lower() for n in ['sanjay', 'shijil', 'parthip', 'sandbox']):
        print(f"P[{i}]: {txt[:120]}")

print("\n=== TABLES WITH NAMES ===")
for ti, t in enumerate(doc.tables):
    found = False
    for ri, r in enumerate(t.rows):
        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
        if any(any(n in c.lower() for n in ['sanjay', 'shijil', 'parthip', 'sandbox']) for c in row_txt):
            if not found:
                print(f"--- Table {ti} ({len(t.rows)} rows x {len(t.columns)} cols) ---")
                found = True
            print(f"  Row {ri}: {' | '.join(row_txt[:6])}")
