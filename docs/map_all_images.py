import docx
import re

d = docx.Document(r"C:\Users\Abinson Babu\.gemini\antigravity\brain\18c1487f-ab35-4cae-a538-c21dfac065a0\.user_uploaded\media_1792596558882_0f1ebc8d.docx")

rels = d.part.rels

for i, p in enumerate(d.paragraphs):
    m = re.findall(r'r:embed="([^"]+)"', p._p.xml)
    if m:
        target = [rels[r].target_ref for r in m if r in rels]
        print(f"=== {target} at paragraph {i} ===")
        for j in range(max(0, i - 4), min(len(d.paragraphs), i + 5)):
            t = d.paragraphs[j].text.strip().encode('ascii', errors='replace').decode()
            if t:
                print(f"  [{j}] {t[:80]}")
